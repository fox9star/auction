from datetime import datetime, timedelta
import openpyxl
import requests
import json
import sys
import time

sys.stdout.reconfigure(encoding='utf-8')

RISK_KEYWORDS = ['유치권', '지분', '법정지상권', '분묘', '가등기', '가처분', '재매각', '위반건축물', '맹지', '토지만매각', '건물만매각']

def compute_100pt_score(usage, discount_rate, yuchal_cnt, risk_str, sido, address, remarks=""):
    # 1. 가격 메리트 (40점 만점 - 세밀화)
    if discount_rate >= 70:
        score_price = 40
    elif discount_rate >= 60:
        score_price = 38
    elif discount_rate >= 50:
        score_price = 35
    elif discount_rate >= 40:
        score_price = 31
    elif discount_rate >= 30:
        score_price = 27
    elif discount_rate >= 20:
        score_price = 22
    elif discount_rate >= 10:
        score_price = 16
    else:
        score_price = 10

    # 2. 권리 안전성 (30점 만점)
    combined_risk_txt = f"{risk_str} {remarks}"
    matched_risks = [rk for rk in RISK_KEYWORDS if rk in combined_risk_txt]
    has_risk = ('주의' in str(risk_str)) or len(matched_risks) > 0
    if not has_risk:
        score_risk = 30
        risk_label = "안전(소멸주의)"
    else:
        score_risk = -15
        rk_tag = ",".join(matched_risks[:2]) if matched_risks else "특수권리확인"
        risk_label = f"주의({rk_tag})"

    # 3. 용도 및 환금성 (20점 만점)
    usg = str(usage or '아파트')
    if any(k in usg for k in ['아파트', '주상복합']):
        score_usg = 20
    elif any(k in usg for k in ['오피스텔', '다세대', '연립', '빌라', '도시형']):
        score_usg = 16
    elif any(k in usg for k in ['단독', '다가구', '상가주택', '주택']):
        score_usg = 14
    elif any(k in usg for k in ['근린', '상가', '점포', '업무', '숙박', '공장']):
        score_usg = 12
    elif any(k in usg for k in ['대지', '택지', '임야', '전', '답']):
        score_usg = 10
    else:
        score_usg = 11

    # 4. 입지 선호도 (10점 만점)
    loc_txt = f"{sido} {address}"
    if any(k in loc_txt for k in ['서울', '유성구', '분당', '수성구', '해운대', '수영구', '과천', '판교', '강남', '서초', '송파', '용산', '마포', '성동']):
        score_loc = 10
    elif any(k in loc_txt for k in ['서구', '수원', '성남', '용인', '고양', '화성', '남동구', '연수구', '부산진', '동래', '달서구', '중구', '천안', '청주', '전주', '광산구', '춘천', '원주', '강릉', '제주시']):
        score_loc = 9
    elif any(k in loc_txt for k in ['경기', '인천', '부산', '대구', '울산', '광주', '대전', '세종']):
        score_loc = 8
    else:
        score_loc = 7

    total_score = max(15, min(100, score_price + score_risk + score_usg + score_loc))

    comments = []
    if not has_risk:
        comments.append("권리분석 안전(특수권리 없음)")
    else:
        comments.append(f"권리상 주의 요망({risk_label})")

    if discount_rate >= 50:
        comments.append(f"감정가 대비 {discount_rate:.1f}% 대폭 할인(유찰 {yuchal_cnt}회)")
    elif discount_rate >= 20:
        comments.append(f"감정가 대비 {discount_rate:.1f}% 저감(유찰 {yuchal_cnt}회)")
    else:
        comments.append("신건 또는 초기 차수 진행")

    if score_usg >= 16:
        comments.append("환금성·실거주 수요 우수 주거용")
    elif score_usg >= 12:
        comments.append("임대수익·실사용 유망 물건")

    return total_score, risk_label, " | ".join(comments)


def fetch_supplement_gangwon_jeju():
    """강원·제주권 물건이 100건 미만일 경우 대법원 서버에서 실시간 추가 수집"""
    session = requests.Session()
    session.headers.update({
        'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36',
        'Accept': 'application/json',
        'Content-Type': 'application/json; charset=UTF-8',
        'Referer': 'https://www.courtauction.go.kr/pgj/index.on'
    })
    extra = []
    for sd_cd, sd_name in [('51', '강원특별자치도'), ('50', '제주특별자치도')]:
        for page_no in range(1, 4):
            payload = {
                "dma_pageInfo": {
                    "pageNo": page_no, "pageSize": 40, "bfPageNo": "", "startRowNo": "",
                    "totalCnt": "", "totalYn": "Y" if page_no == 1 else "N", "groupTotalCount": ""
                },
                "dma_srchGdsDtlSrchInfo": {
                    "rletDspslSpcCondCd": "", "bidDvsCd": "", "mvprpRletDvsCd": "00031R",
                    "cortAuctnSrchCondCd": "0004601", "rprsAdongSdCd": sd_cd,
                    "rprsAdongSggCd": "", "rprsAdongEmdCd": "", "rdnmSdCd": "", "rdnmSggCd": "", "rdnmNo": "",
                    "mvprpDspslPlcAdongSdCd": "", "mvprpDspslPlcAdongSggCd": "", "mvprpDspslPlcAdongEmdCd": "",
                    "rdDspslPlcAdongSdCd": "", "rdDspslPlcAdongSggCd": "", "rdDspslPlcAdongEmdCd": "",
                    "cortOfcCd": "", "jdbnCd": "", "execrOfcDvsCd": "",
                    "lclDspslGdsLstUsgCd": "", "mclDspslGdsLstUsgCd": "", "sclDspslGdsLstUsgCd": "",
                    "cortAuctnMbrsId": "", "aeeEvlAmtMin": "", "aeeEvlAmtMax": "",
                    "lwsDspslPrcRateMin": "", "lwsDspslPrcRateMax": "",
                    "flbdNcntMin": "", "flbdNcntMax": "",
                    "objctArDtsMin": "", "objctArDtsMax": "",
                    "mvprpArtclKndCd": "", "mvprpArtclNm": "", "mvprpAtchmPlcTypCd": "",
                    "notifyLoc": "on", "lafjOrderBy": "", "pgmId": "PGJ151F01", "csNo": "",
                    "cortStDvs": "2", "statNum": 1,
                    "bidBgngYmd": datetime.now().strftime("%Y%m%d"),
                    "bidEndYmd": (datetime.now() + timedelta(days=45)).strftime("%Y%m%d"),
                    "dspslDxdyYmd": "", "fstDspslHm": "", "scdDspslHm": "", "thrdDspslHm": "", "fothDspslHm": "",
                    "dspslPlcNm": "", "lwsDspslPrcMin": "", "lwsDspslPrcMax": "",
                    "grbxTypCd": "", "gdsVendNm": "", "fuelKndCd": "", "carMdyrMax": "", "carMdyrMin": "", "carMdlNm": ""
                }
            }
            try:
                res = session.post('https://www.courtauction.go.kr/pgj/pgjsearch/searchControllerMain.on', json=payload, timeout=10)
                if res.status_code == 200:
                    dlt = res.json().get('data', {}).get('dlt_srchResult', [])
                    for it in dlt:
                        app = int(it.get('gamevalAmt') or 0)
                        min_p = int(it.get('notifyMinmaePrice1') or it.get('minmaePrice') or 0)
                        yc = int(it.get('yuchalCnt') or 0)
                        if app > 0 and min_p > 0:
                            rat = round((min_p / app) * 100, 1)
                            disc = round(100 - rat, 1)
                        else:
                            rat = 100.0
                            disc = 0.0
                        md = it.get('maeGiil', '')
                        md_fmt = f"{md[:4]}-{md[4:6]}-{md[6:8]}" if len(md) == 8 else md
                        usg = it.get('dspslUsgNm', '주택')
                        addr = it.get('printSt', '')
                        bigo = it.get('mulBigo') or ''
                        sc, rk_lbl, cmt = compute_100pt_score(usg, disc, yc, "", sd_name, addr, bigo)
                        extra.append({
                            'region': '강원·제주권',
                            'sido': sd_name,
                            'sigu': it.get('hjguSigu') or sd_name,
                            'court': it.get('jiwonNm', ''),
                            'dept': it.get('jpDeptNm', '경매계'),
                            'case_no': str(it.get('srnSaNo', '')),
                            'item_no': str(it.get('maemulSer', '1')),
                            'usage': usg,
                            'address': addr,
                            'spec': it.get('convAddr', ''),
                            'appraisal': app,
                            'min_price': min_p,
                            'ratio': f"{rat}%",
                            'discount': f"{disc}%",
                            'discount_val': disc,
                            'yuchal': f"{yc}회",
                            'date': md_fmt,
                            'risk': rk_lbl,
                            'comment': cmt,
                            'score': sc
                        })
            except Exception as e:
                print(f"Supplement fetch warning ({sd_name} p{page_no}): {e}")
    return extra


# 1. 대전 전체 목록 파싱 -> 대전 TOP 100 구성
wb_dj = openpyxl.load_workbook(r'C:\Users\wood\Desktop\aiwork\auction\대전광역시_경매물건_TOP50_추천.xlsx', data_only=True)
ws_dj_all = wb_dj['대전_전체물건_전수목록']

dj_all_raw = []
for r in ws_dj_all.iter_rows(min_row=5, values_only=True):
    if r[0] is None:
        continue
    # 전체물건 컬럼: 번호(0), 종합점수(1), 자치구(2), 사건번호(3), 물건번호(4), 용도(5), 소재지(6), 내역(7), 감정평가액(8), 최저매각가격(9), 최저가율(10), 할인율(11), 유찰횟수(12), 매각기일(13), 매각시간(14), 매각장소(15), 담당계(16), 연락처(17), 권리구분(18), 비고(19)
    score_val = int(r[1] or 80)
    disc_str = str(r[11] or '0%')
    disc_val = float(disc_str.replace('%', '').strip() or 0)
    yc_str = str(r[12] or '0회')
    yc_val = int(yc_str.replace('회', '').strip() or 0)
    usg = str(r[5] or '아파트')
    sigu = str(r[2] or '대전')
    addr = str(r[6] or '')
    risk_s = str(r[18] or '안전(소멸주의)')
    bigo_s = str(r[19] or '')
    _, _, cmt = compute_100pt_score(usg, disc_val, yc_val, risk_s, '대전광역시 ' + sigu, addr, bigo_s)
    dj_all_raw.append({
        'score': score_val,
        'region': '충청권',
        'sido': '대전광역시',
        'sigu': sigu,
        'court': '대전지방법원',
        'dept': str(r[16] or '경매계'),
        'case_no': str(r[3]),
        'item_no': str(r[4]),
        'usage': usg,
        'address': addr,
        'spec': str(r[7] or ''),
        'appraisal': int(r[8] or 0),
        'min_price': int(r[9] or 0),
        'ratio': str(r[10]),
        'discount': disc_str if '%' in disc_str else f"{disc_val:.1f}%",
        'discount_val': disc_val,
        'yuchal': yc_str if '회' in yc_str else f"{yc_val}회",
        'date': str(r[13] or ''),
        'risk': risk_s,
        'comment': cmt
    })

dj_total_cnt = len(dj_all_raw)
dj_all_raw.sort(key=lambda x: (x['score'], x['discount_val'], -x['min_price']), reverse=True)

dj_top100 = []
for idx, it in enumerate(dj_all_raw[:100], start=1):
    item_copy = dict(it)
    item_copy['id'] = f"DJ_{idx}"
    item_copy['rank'] = idx
    item_copy['rank_label'] = f"대전 {idx}위"
    item_copy['region_key'] = '대전광역시'
    dj_top100.append(item_copy)

# 2. 전국 전체 목록 파싱 -> 권역별 TOP 100 구성
wb_nat = openpyxl.load_workbook(r'C:\Users\wood\Desktop\aiwork\auction\대한민국법원_전국경매물건_정리_및_권역별추천.xlsx', data_only=True)
ws_all = wb_nat['전국_경매물건_전체목록']

nat_all_raw = []
for r in ws_all.iter_rows(min_row=5, values_only=True):
    if r[0] is None:
        continue
    # 번호(0), 권역(1), 시도(2), 관할법원(3), 담당계(4), 사건번호(5), 물건번호(6), 용도(7), 주소(8), 내역(9), 감정가(10), 최저가(11), 최저가율(12), 할인율(13), 유찰횟수(14), 매각기일(15), 매각시간(16), 매각장소(17), 권리구분(18), 비고(19)
    reg = str(r[1] or '')
    sido = str(r[2] or '')
    usg = str(r[7] or '아파트')
    addr = str(r[8] or '')
    disc_str = str(r[13] or '0%')
    disc_val = float(disc_str.replace('%', '').strip() or 0)
    yc_str = str(r[14] or '0회')
    yc_val = int(yc_str.replace('회', '').strip() or 0)
    risk_s = str(r[18] or '일반(안전)')
    bigo_s = str(r[19] or '')
    sc, rk_lbl, cmt = compute_100pt_score(usg, disc_val, yc_val, risk_s, sido, addr, bigo_s)
    nat_all_raw.append({
        'region': reg,
        'sido': sido,
        'sigu': sido,
        'court': str(r[3] or ''),
        'dept': str(r[4] or '경매계'),
        'case_no': str(r[5]),
        'item_no': str(r[6]),
        'usage': usg,
        'address': addr,
        'spec': str(r[9] or ''),
        'appraisal': int(r[10] or 0),
        'min_price': int(r[11] or 0),
        'ratio': str(r[12]),
        'discount': disc_str if '%' in disc_str else f"{disc_val:.1f}%",
        'discount_val': disc_val,
        'yuchal': yc_str if '회' in yc_str else f"{yc_val}회",
        'date': str(r[15] or ''),
        'risk': rk_lbl,
        'comment': cmt,
        'score': sc
    })

# 강원·제주권이 100건 미만이면 실시간 추가 수집하여 100건 이상 확보
gw_curr = [x for x in nat_all_raw if '강원' in x['region'] or '제주' in x['region']]
if len(gw_curr) < 100:
    print(f"강원·제주권 현재 {len(gw_curr)}건 -> 대법원 실시간 추가 수집으로 TOP 100 보강 중...")
    extra_gw = fetch_supplement_gangwon_jeju()
    seen_gw = {(x['court'], x['case_no'], x['item_no']) for x in gw_curr}
    for ex in extra_gw:
        k = (ex['court'], ex['case_no'], ex['item_no'])
        if k not in seen_gw:
            seen_gw.add(k)
            nat_all_raw.append(ex)

nat_total_cnt = len(nat_all_raw)

# 권역별 그룹화 및 TOP 100 추출
# 1) 수도권 TOP 100
# 2) 영남권 통합 TOP 100
# 3) 영남권(부울경) TOP 100
# 4) 영남권(대경권) TOP 100
# 5) 충청권 TOP 100 (충청권 전국 데이터 + 대전 데이터 통합 중복제거 후 상위 100선)
# 6) 호남권 TOP 100
# 7) 강원·제주권 TOP 100

def build_region_top100(items_pool, region_key, prefix_id, rank_prefix, limit=100):
    dedup = {}
    for x in items_pool:
        k = f"{x['court']}_{x['case_no']}_{x['item_no']}"
        if k not in dedup or x['score'] > dedup[k]['score']:
            dedup[k] = x
    sorted_pool = sorted(dedup.values(), key=lambda x: (x['score'], x['discount_val'], -x['min_price']), reverse=True)
    result = []
    for idx, it in enumerate(sorted_pool[:limit], start=1):
        c = dict(it)
        c['id'] = f"{prefix_id}_{idx}"
        c['rank'] = idx
        c['rank_label'] = f"{rank_prefix} {idx}위"
        c['region_key'] = region_key
        result.append(c)
    return result

sudo_pool = [x for x in nat_all_raw if x['region'] == '수도권']
yeongnam_all_pool = [x for x in nat_all_raw if '영남권' in x['region']]
pk_pool = [x for x in nat_all_raw if x['region'] == '영남권(부울경)']
tk_pool = [x for x in nat_all_raw if x['region'] == '영남권(대경권)']
chung_pool = [x for x in nat_all_raw if x['region'] == '충청권'] + dj_all_raw
honam_pool = [x for x in nat_all_raw if x['region'] == '호남권']
gw_pool = [x for x in nat_all_raw if '강원' in x['region'] or '제주' in x['region']]

regional_top100 = {
    '대전광역시': dj_top100,
    '수도권': build_region_top100(sudo_pool, '수도권', 'SUDO', '수도권', 100),
    '영남권': build_region_top100(yeongnam_all_pool, '영남권', 'YN', '영남권', 100),
    '영남권(부울경)': build_region_top100(pk_pool, '영남권(부울경)', 'PK', '부울경', 100),
    '영남권(대경권)': build_region_top100(tk_pool, '영남권(대경권)', 'TK', '대경권', 100),
    '충청권': build_region_top100(chung_pool, '충청권', 'CC', '충청권', 100),
    '호남권': build_region_top100(honam_pool, '호남권', 'HN', '호남권', 100),
    '강원': build_region_top100(gw_pool, '강원', 'GW', '강원·제주', 100),
}

for k, v in regional_top100.items():
    print(f" - [{k}] TOP {len(v)} 구축 완료 (1위 {v[0]['score']}점 ~ {len(v)}위 {v[-1]['score']}점)")

# 전국 베스트 (권역별 대표 상위 3개씩 18선)
nat_best = []
best_idx = 1
for reg_k in ['수도권', '영남권(부울경)', '영남권(대경권)', '충청권', '호남권', '강원']:
    for it in regional_top100[reg_k][:3]:
        c = dict(it)
        c['id'] = f"NAT_{best_idx}"
        c['rank'] = best_idx
        c['rank_label'] = f"전국 베스트 {best_idx}위"
        nat_best.append(c)
        best_idx += 1

all_combined = []
for v in regional_top100.values():
    all_combined.extend(v)

max_disc = 0.0
for it in all_combined:
    try:
        v = float(str(it.get('discount', '0')).replace('%', '').strip())
        if v > max_disc:
            max_disc = v
    except Exception:
        pass

data = {
    'meta': {
        'refreshed_at': datetime.now().strftime('%Y-%m-%d %H:%M:%S'),
        'nationwide_total': nat_total_cnt,
        'daejeon_total': dj_total_cnt,
        'daejeon_top_count': len(dj_top100),
        'region_top_count': 100,
        'max_discount': f"{max_disc:.1f}%",
        'region_counts': {k: len(v) for k, v in regional_top100.items()}
    },
    'daejeon_top50': dj_top100,
    'regional_top100': regional_top100,
    'nationwide_best': nat_best,
    'nationwide_more': []
}

out_file = r'C:\Users\wood\Desktop\aiwork\auction\portal_data.json'
try:
    from enrich_engine import enrich_all_portal_data
    data = enrich_all_portal_data(data)
except Exception as e:
    print('Enrichment warning:', e)

with open(out_file, 'w', encoding='utf-8') as f:
    json.dump(data, f, ensure_ascii=False, indent=2)

print(f"\nportal_data.json 저장 완료! 전국 총 {nat_total_cnt}건, 대전 총 {dj_total_cnt}건, 각 권역별 TOP 100 완료.")
