import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

sys.stdout.reconfigure(encoding='utf-8')

# 16:9 와이드스크린 설정
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

blank_layout = prs.slide_layouts[6]

# 색상 팔레트 정의
COLOR_NAVY = RGBColor(26, 54, 93)      # #1A365D (주요 배경/제목)
COLOR_BLUE = RGBColor(43, 108, 176)    # #2B6CB0 (강조색)
COLOR_TEAL = RGBColor(49, 151, 149)    # #319795 (서브 강조)
COLOR_GRAY_BG = RGBColor(247, 250, 252)# #F7FAFC (슬라이드 배경)
COLOR_CARD_BG = RGBColor(255, 255, 255)# #FFFFFF
COLOR_TEXT_MAIN = RGBColor(45, 55, 72) # #2D3748
COLOR_TEXT_MUTED = RGBColor(113, 128, 150) # #718096
COLOR_WHITE = RGBColor(255, 255, 255)
COLOR_RED = RGBColor(197, 48, 48)      # 할인/주의
COLOR_GREEN = RGBColor(47, 133, 90)    # 안전/추천
COLOR_LINE = RGBColor(226, 232, 240)

IMG_DIR = r"C:\Users\wood\Desktop\aiwork\auction\images"
img_banner = os.path.join(IMG_DIR, "hero_auction_banner.jpg")
img_extinction = os.path.join(IMG_DIR, "infographic_extinction.png")
img_nationwide = os.path.join(IMG_DIR, "chart_nationwide.png")
img_daejeon = os.path.join(IMG_DIR, "chart_daejeon_top50.png")

def set_slide_background(slide, color):
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg.fill.solid()
    bg.fill.fore_color.rgb = color
    bg.line.fill.background()
    return bg

def add_header(slide, title_text, category="대한민국 법원경매 종합 분석"):
    cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(10), Inches(0.4))
    tf_c = cat_box.text_frame
    p_c = tf_c.paragraphs[0]
    p_c.text = category.upper()
    p_c.font.size = Pt(11)
    p_c.font.bold = True
    p_c.font.color.rgb = COLOR_TEAL

    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(11.5), Inches(0.8))
    tf_t = title_box.text_frame
    p_t = tf_t.paragraphs[0]
    p_t.text = title_text
    p_t.font.size = Pt(22)
    p_t.font.bold = True
    p_t.font.color.rgb = COLOR_NAVY

    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.5), Inches(11.733), Inches(0.02))
    line.fill.solid()
    line.fill.fore_color.rgb = COLOR_LINE
    line.line.fill.background()

# ==============================================================================
# SLIDE 1: 표지 (Title Slide with Hero Banner Image)
# ==============================================================================
slide1 = prs.slides.add_slide(blank_layout)
set_slide_background(slide1, COLOR_NAVY)

# 좌측 텍스트 영역 + 우측 고화질 배너 이미지 배치
if os.path.exists(img_banner):
    # 우측 배너 이미지 삽입
    slide1.shapes.add_picture(img_banner, Inches(6.8), Inches(1.2), Inches(5.8), Inches(5.2))
    
    # 우측 배너 테두리 프레임
    frame = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(6.8), Inches(1.2), Inches(5.8), Inches(5.2))
    frame.fill.background()
    frame.line.color.rgb = COLOR_TEAL
    frame.line.width = Pt(2)

# 좌측 타이틀
tb_t = slide1.shapes.add_textbox(Inches(0.8), Inches(1.4), Inches(5.8), Inches(3.2))
tf = tb_t.text_frame
tf.word_wrap = True

p0 = tf.paragraphs[0]
p0.text = "REAL ESTATE AUCTION"
p0.font.size = Pt(12)
p0.font.bold = True
p0.font.color.rgb = COLOR_TEAL

p1 = tf.add_paragraph()
p1.text = "대한민국 법원\n부동산 경매 종합 분석"
p1.font.size = Pt(32)
p1.font.bold = True
p1.font.color.rgb = COLOR_WHITE
p1.space_before = Pt(8)

p2 = tf.add_paragraph()
p2.text = "실시간 전국 데이터 분석·소멸주의 가이드 및 대전광역시 TOP 50 가치평가 모델"
p2.font.size = Pt(13)
p2.font.color.rgb = RGBColor(203, 213, 225)
p2.space_before = Pt(14)

# 좌측 하단 메타 박스
meta_box = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(5.1), Inches(5.6), Inches(1.3))
meta_box.fill.solid()
meta_box.fill.fore_color.rgb = RGBColor(30, 64, 110)
meta_box.line.color.rgb = RGBColor(43, 108, 176)

tb_m = slide1.shapes.add_textbox(Inches(0.95), Inches(5.2), Inches(5.3), Inches(1.1))
tf_m = tb_m.text_frame
p_m = tf_m.paragraphs[0]
p_m.text = "• 데이터 출처: 대한민국 법원경매정보 (courtauction.go.kr)\n• 수집 규모: 전국 1,287건 / 대전 522건 전수 데이터 분석\n• 산출물: 종합 분석 보고서 및 엑셀 데이터시트 2종 탑재"
p_m.font.size = Pt(10.5)
p_m.font.color.rgb = RGBColor(226, 232, 240)

# ==============================================================================
# SLIDE 2: 목차 (Table of Contents)
# ==============================================================================
slide2 = prs.slides.add_slide(blank_layout)
set_slide_background(slide2, COLOR_GRAY_BG)
add_header(slide2, "보고서 핵심 구성 및 목차", "Table of Contents")

toc_items = [
    ("01", "부산동경매 정의 및 법원 체계", "부산지방법원 동부지원 관할 구역(해운대·수영·남구·기장) 및 입찰법정 특징"),
    ("02", "전국 60개 법원 입찰법정 및 준비물", "전국 18개 본원 + 42개 지원 현장 응찰 체계 및 필수 준비 서류"),
    ("03", "경매 물건 선별 5대 핵심 지표", "가격 괴리·안전마진·소멸주의·환금성·명도 관리비 5단계 실전 공식"),
    ("04", "소멸주의(消滅主義) 권리분석 인포그래픽", "말소기준권리 5총사, 인수주의와의 비교 및 구매자 3대 제도적 이점"),
    ("05", "전국 권역별 경매 데이터 통계 & 추천", "전국 1,287건 점유율/할인율 차트 분석 및 권역별 가성비 추천 매물"),
    ("06", "대전광역시 경매 분석 & TOP 50", "대전 522건 전수 수집, 자치구별 시각화 및 100점 만점 베스트 매물")
]

for idx, (num, title, desc) in enumerate(toc_items):
    col = idx % 2
    row = idx // 2
    x = Inches(0.8 + col * 5.95)
    y = Inches(1.8 + row * 1.65)
    
    card = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(5.75), Inches(1.4))
    card.fill.solid()
    card.fill.fore_color.rgb = COLOR_CARD_BG
    card.line.color.rgb = COLOR_LINE
    
    num_box = slide2.shapes.add_textbox(x + Inches(0.2), y + Inches(0.2), Inches(0.9), Inches(1.0))
    p_n = num_box.text_frame.paragraphs[0]
    p_n.text = num
    p_n.font.size = Pt(26)
    p_n.font.bold = True
    p_n.font.color.rgb = COLOR_BLUE
    
    txt_box = slide2.shapes.add_textbox(x + Inches(1.1), y + Inches(0.2), Inches(4.4), Inches(1.0))
    tf_txt = txt_box.text_frame
    tf_txt.word_wrap = True
    p_t = tf_txt.paragraphs[0]
    p_t.text = title
    p_t.font.size = Pt(14)
    p_t.font.bold = True
    p_t.font.color.rgb = COLOR_NAVY
    
    p_d = tf_txt.add_paragraph()
    p_d.text = desc
    p_d.font.size = Pt(10)
    p_d.font.color.rgb = COLOR_TEXT_MUTED
    p_d.space_before = Pt(4)

# ==============================================================================
# SLIDE 3: 부산동경매의 정의 및 특징
# ==============================================================================
slide3 = prs.slides.add_slide(blank_layout)
set_slide_background(slide3, COLOR_GRAY_BG)
add_header(slide3, "부산동경매(부산지방법원 동부지원)의 핵심 특징", "Section 01")

card1 = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(5.7), Inches(5.0))
card1.fill.solid()
card1.fill.fore_color.rgb = COLOR_CARD_BG
card1.line.color.rgb = COLOR_LINE

tb1 = slide3.shapes.add_textbox(Inches(1.1), Inches(2.0), Inches(5.1), Inches(4.6))
tf1 = tb1.text_frame
tf1.word_wrap = True

p = tf1.paragraphs[0]
p.text = "🏛 부산지방법원 동부지원 기본 정보"
p.font.size = Pt(16)
p.font.bold = True
p.font.color.rgb = COLOR_NAVY

p = tf1.add_paragraph()
p.text = "• 약칭: 부산동 / 부산동부 경매\n• 소재지: 부산광역시 해운대구 재반로112번길 20 (재송동)\n• 관할 구역 (4개 시·구·군):\n   - 해운대구 (센텀시티, 마린시티, 우동, 중동, 좌동)\n   - 수영구 (남천동, 광안동, 민락동, 망미동)\n   - 남구 (대연동, 용호동, 문현동 등 재개발 구역)\n   - 기장군 (정관·일광신도시, 기장읍, 장안읍, 철마면)\n• 입찰 시간:\n   - 입찰서 접수: 오전 10:00 ~ 오전 11:20\n   - 개찰 시작: 오전 11:30 순차 진행"
p.font.size = Pt(11.5)
p.font.color.rgb = COLOR_TEXT_MAIN
p.space_before = Pt(10)

card2 = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.8), Inches(5.7), Inches(5.0))
card2.fill.solid()
card2.fill.fore_color.rgb = COLOR_CARD_BG
card2.line.color.rgb = COLOR_LINE

tb2 = slide3.shapes.add_textbox(Inches(7.1), Inches(2.0), Inches(5.1), Inches(4.6))
tf2 = tb2.text_frame
tf2.word_wrap = True

p = tf2.paragraphs[0]
p.text = "📈 시장 특징 및 관할 주의사항"
p.font.size = Pt(16)
p.font.bold = True
p.font.color.rgb = COLOR_NAVY

p = tf2.add_paragraph()
p.text = "• 부산 최고의 주거 선호지 밀집:\n   - 해운대·수영·남구는 부산 아파트 경매의 최대 격전지\n   - 낙찰가율 및 입찰 경쟁률이 부산 본원 대비 높음\n\n• 기장군 도농복합·개발 호재지:\n   - 오시리아 관광단지 배후 상가 및 택지지구 아파트\n   - 토지·임야·공장 등 투자용 물건 다양\n\n• ⚠️ 관할 주의: '부산 동구' 지역 물건\n   - 초량동, 범일동, 수정동, 좌천동 등 '부산 동구' 물건은\n     동부지원이 아니라 '부산지방법원 본원'(연제구 거제동) 관할!"
p.font.size = Pt(11.5)
p.font.color.rgb = COLOR_TEXT_MAIN
p.space_before = Pt(10)

# ==============================================================================
# SLIDE 4: 전국 60개 법원 입찰법정 체계
# ==============================================================================
slide4 = prs.slides.add_slide(blank_layout)
set_slide_background(slide4, COLOR_GRAY_BG)
add_header(slide4, "전국 60개 법원 경매 입찰법정 및 응찰 절차", "Section 02")

bar = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(11.733), Inches(0.8))
bar.fill.solid()
bar.fill.fore_color.rgb = RGBColor(235, 248, 255)
bar.line.color.rgb = COLOR_BLUE

tb_b = slide4.shapes.add_textbox(Inches(1.0), Inches(1.85), Inches(11.3), Inches(0.7))
tf_b = tb_b.text_frame
p = tf_b.paragraphs[0]
p.text = "📌 현장 입찰 원칙: 대한민국 법원경매는 온라인 입찰이 불가하며, 관할 법원 입찰법정에 직접 출석해야 합니다. (총 60개소)"
p.font.size = Pt(12)
p.font.bold = True
p.font.color.rgb = COLOR_NAVY

regions_data = [
    ("수도권 (16개소)", "• 서울 5개 본원 (중앙, 동부, 서부, 남부, 북부)\n• 인천 본원, 부천지원\n• 의정부 본원, 고양지원, 남양주지원\n• 수원 본원, 성남, 여주, 평택, 안산, 안양지원"),
    ("영남권 (19개소)", "• 부산 본원, 부산동부지원, 부산서부지원\n• 울산 본원\n• 창원 본원, 마산, 진주, 통영, 밀양, 거창지원\n• 대구 본원, 대구서부, 안동, 경주, 포항, 김천, 상주, 의성, 영덕지원"),
    ("충청·강원권 (15개소)", "• 대전 본원, 홍성, 논산, 천안, 공주, 서산지원\n• 청주 본원, 충주, 제천, 영동지원\n• 춘천 본원, 강릉, 원주, 속초, 영월지원"),
    ("호남·제주권 (10개소)", "• 광주 본원, 목포, 장흥, 순천, 해남지원\n• 전주 본원, 군산, 정읍, 남원지원\n• 제주지방법원 본원")
]

for idx, (r_name, r_desc) in enumerate(regions_data):
    x = Inches(0.8 + idx * 2.98)
    y = Inches(2.8)
    card = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(2.8), Inches(4.1))
    card.fill.solid()
    card.fill.fore_color.rgb = COLOR_CARD_BG
    card.line.color.rgb = COLOR_LINE
    
    tb = slide4.shapes.add_textbox(x + Inches(0.15), y + Inches(0.2), Inches(2.5), Inches(3.7))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = r_name
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = COLOR_BLUE
    
    p = tf.add_paragraph()
    p.text = r_desc
    p.font.size = Pt(10)
    p.font.color.rgb = COLOR_TEXT_MAIN
    p.space_before = Pt(8)

# ==============================================================================
# SLIDE 5: 경매 물건 선별 5대 지표
# ==============================================================================
slide5 = prs.slides.add_slide(blank_layout)
set_slide_background(slide5, COLOR_GRAY_BG)
add_header(slide5, "실전 경매 물건 선별 5대 지표 & 체크리스트", "Section 03")

steps_data = [
    ("1단계", "가격 및 수익성", "감정평가액 맹신 금지\n현재 급매가 대비 안전마진 확인\n유찰 저감률(20~30%) 분석"),
    ("2단계", "권리분석 안전성", "말소기준권리 이후 소멸 여부\n선순위 대항력 임차인 인수액 0원\n특수권리(유치권/지분) 필터링"),
    ("3단계", "시장 경쟁도·환금성", "최근 3~6개월 매매/전세 실거래량\n평균 입찰 경쟁자 수 확인\n낙찰 후 빠른 엑시트 가능성"),
    ("4단계", "명도 및 숨은 비용", "점유자 성향 및 전액 배당 여부\n인도명령 및 강제집행 비용 고려\n공용부분 미납 관리비 사전 확인"),
    ("5단계", "최종 입찰가 결정", "급매 시세 - 총 취득원가\n(취득세+명도비+수리비+대출이자)\n= 나의 목표 입찰가 도출")
]

for idx, (st_num, st_title, st_desc) in enumerate(steps_data):
    x = Inches(0.8 + idx * 2.38)
    y = Inches(1.8)
    card = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(2.25), Inches(5.0))
    card.fill.solid()
    card.fill.fore_color.rgb = COLOR_CARD_BG
    card.line.color.rgb = COLOR_LINE
    
    st_box = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x + Inches(0.2), y + Inches(0.2), Inches(1.85), Inches(0.45))
    st_box.fill.solid()
    st_box.fill.fore_color.rgb = COLOR_BLUE
    st_box.line.fill.background()
    p = st_box.text_frame.paragraphs[0]
    p.text = st_num
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = COLOR_WHITE
    p.alignment = PP_ALIGN.CENTER
    
    tb = slide5.shapes.add_textbox(x + Inches(0.15), y + Inches(0.8), Inches(1.95), Inches(4.0))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = st_title
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = COLOR_NAVY
    p.alignment = PP_ALIGN.CENTER
    
    p = tf.add_paragraph()
    p.text = st_desc
    p.font.size = Pt(10)
    p.font.color.rgb = COLOR_TEXT_MAIN
    p.space_before = Pt(10)

# ==============================================================================
# SLIDE 6: 소멸주의 메커니즘 인포그래픽 (IMAGE INSERTED)
# ==============================================================================
slide6 = prs.slides.add_slide(blank_layout)
set_slide_background(slide6, COLOR_GRAY_BG)
add_header(slide6, "소멸주의(消滅主義) 권리 소멸 메커니즘 인포그래픽", "Section 04")

if os.path.exists(img_extinction):
    # 상단 인포그래픽 이미지 삽입
    slide6.shapes.add_picture(img_extinction, Inches(0.8), Inches(1.8), Inches(11.733), Inches(4.2))
    
    # 하단 요약 설명 바
    summary_box = slide6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(6.15), Inches(11.733), Inches(0.85))
    summary_box.fill.solid()
    summary_box.fill.fore_color.rgb = RGBColor(240, 253, 244)
    summary_box.line.color.rgb = COLOR_GREEN
    
    tb_s = slide6.shapes.add_textbox(Inches(1.0), Inches(6.2), Inches(11.3), Inches(0.75))
    tf_s = tb_s.text_frame
    p = tf_s.paragraphs[0]
    p.text = "💡 핵심 원리: 낙찰자가 잔금을 완납하면 말소기준권리(근저당·가압류 등)를 포함한 후순위 채권은 법원 촉탁으로 100% 소멸(말소)됩니다. 단, 말소기준보다 빠른 '대항력 있는 선순위 임차인의 미배당 보증금'과 '유치권'은 낙찰자가 인수하므로 사전 권리분석이 필수입니다."
    p.font.size = Pt(10)
    p.font.color.rgb = COLOR_GREEN
    p.font.bold = True

# ==============================================================================
# SLIDE 7: 전국 권역별 경매 데이터 통계 차트 (IMAGE INSERTED)
# ==============================================================================
slide7 = prs.slides.add_slide(blank_layout)
set_slide_background(slide7, COLOR_GRAY_BG)
add_header(slide7, "전국 법원경매 물건 권역별 점유율 및 할인율 분석", "Section 05")

if os.path.exists(img_nationwide):
    slide7.shapes.add_picture(img_nationwide, Inches(0.8), Inches(1.8), Inches(11.733), Inches(4.2))
    
    info_box = slide7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(6.15), Inches(11.733), Inches(0.85))
    info_box.fill.solid()
    info_box.fill.fore_color.rgb = RGBColor(235, 248, 255)
    info_box.line.color.rgb = COLOR_BLUE
    
    tb_i = slide7.shapes.add_textbox(Inches(1.0), Inches(6.2), Inches(11.3), Inches(0.75))
    tf_i = tb_i.text_frame
    p = tf_i.paragraphs[0]
    p.text = "📊 데이터 분석 요약: 전국 1,287건 중 수도권(484건, 37.6%)과 영남권(337건, 26.2%)이 전체의 64%를 차지합니다. 3회 유찰 시 모든 권역에서 최대 65.7%의 파격적 할인율(최저가율 34.3%)이 형성되어 감정가 대비 압도적인 가격 메리트를 제공합니다."
    p.font.size = Pt(10)
    p.font.color.rgb = COLOR_NAVY
    p.font.bold = True

# ==============================================================================
# SLIDE 8: 전국 권역별 베스트 추천 물건 표
# ==============================================================================
slide8 = prs.slides.add_slide(blank_layout)
set_slide_background(slide8, COLOR_GRAY_BG)
add_header(slide8, "전국 권역별 최고 가성비 추천 물건 현황", "Section 05 - Best Picks")

table_shape = slide8.shapes.add_table(7, 7, Inches(0.8), Inches(1.8), Inches(11.733), Inches(5.0))
table = table_shape.table

table.columns[0].width = Inches(1.4)
table.columns[1].width = Inches(1.4)
table.columns[2].width = Inches(1.6)
table.columns[3].width = Inches(3.5)
table.columns[4].width = Inches(1.3)
table.columns[5].width = Inches(1.3)
table.columns[6].width = Inches(1.2)

headers = ["권역", "관할법원", "사건번호", "소재지 및 아파트명", "감정평가액", "최저매각가격", "할인율"]
for col_idx, h in enumerate(headers):
    cell = table.cell(0, col_idx)
    cell.text = h
    cell.fill.solid()
    cell.fill.fore_color.rgb = COLOR_BLUE
    p = cell.text_frame.paragraphs[0]
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = COLOR_WHITE
    p.alignment = PP_ALIGN.CENTER

rec_data = [
    ("수도권", "부천지원", "2024타경5702", "경기 부천시 소사본동 삼신아파트 1동 3층", "2억 6,700만", "9,158만 원", "-65.7% (3회)"),
    ("영남(부울경)", "부산동부", "2025타경462", "부산 수영구 광안동 성우이린타워 9층", "2억 2,800만", "7,820만 원", "-65.7% (3회)"),
    ("영남(부울경)", "부산동부", "2025타경634", "부산 수영구 망미동 망미한신 103동 2층", "5억 7,500만", "1억 9,722만 원", "-65.7% (3회)"),
    ("영남(대경권)", "대구지법", "2025타경10173", "대구 동구 신서동 명가클래식 11층", "1억 9,100만", "6,551만 원", "-65.7% (3회)"),
    ("충청권", "대전지법", "2025타경503447", "대전 대덕구 대화동 세원풍남 5동 5층", "1억 6,600만", "5,693만 원", "-65.7% (3회)"),
    ("강원권", "속초지원", "2025타경10411", "강원 양양군 양양읍 양양서문주공 102동", "1억 5,000만", "5,145만 원", "-65.7% (3회)")
]

for row_idx, r in enumerate(rec_data, start=1):
    for col_idx, val in enumerate(r):
        cell = table.cell(row_idx, col_idx)
        cell.text = val
        cell.fill.solid()
        cell.fill.fore_color.rgb = COLOR_CARD_BG if row_idx % 2 == 1 else RGBColor(241, 245, 249)
        p = cell.text_frame.paragraphs[0]
        p.font.size = Pt(10)
        p.font.color.rgb = COLOR_TEXT_MAIN
        if col_idx in [0, 1, 2, 6]:
            p.alignment = PP_ALIGN.CENTER
        elif col_idx in [4, 5]:
            p.alignment = PP_ALIGN.RIGHT
            if col_idx == 5:
                p.font.bold = True

# ==============================================================================
# SLIDE 9: 대전광역시 경매 데이터 시각화 차트 (IMAGE INSERTED)
# ==============================================================================
slide9 = prs.slides.add_slide(blank_layout)
set_slide_background(slide9, COLOR_GRAY_BG)
add_header(slide9, "대전광역시 TOP 50 자치구별 분포 및 점수 시각화", "Section 06")

if os.path.exists(img_daejeon):
    slide9.shapes.add_picture(img_daejeon, Inches(0.8), Inches(1.8), Inches(11.733), Inches(4.2))
    
    info_box9 = slide9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(6.15), Inches(11.733), Inches(0.85))
    info_box9.fill.solid()
    info_box9.fill.fore_color.rgb = RGBColor(255, 245, 245)
    info_box9.line.color.rgb = COLOR_RED
    
    tb_i9 = slide9.shapes.add_textbox(Inches(1.0), Inches(6.2), Inches(11.3), Inches(0.75))
    tf_i9 = tb_i9.text_frame
    p = tf_i9.paragraphs[0]
    p.text = "🎯 대전 TOP 50 분석: 주거 및 상업 선호도가 가장 높은 '유성구(22건, 44%)'와 '서구(17건, 34%)'가 상위 78%를 독점하고 있습니다. 100점 만점 기준 95점 이상 물건이 대거 포진해 있으며 감정가 대비 최고 83%~98%까지 하락한 초특가 매물이 존재합니다."
    p.font.size = Pt(10)
    p.font.color.rgb = COLOR_RED
    p.font.bold = True

# ==============================================================================
# SLIDE 10: 대전광역시 TOP 10 추천 물건 표
# ==============================================================================
slide10 = prs.slides.add_slide(blank_layout)
set_slide_background(slide10, COLOR_GRAY_BG)
add_header(slide10, "대전광역시 경매 베스트 TOP 10 추천 물건", "Section 06 - Daejeon TOP 10")

table_shape10 = slide10.shapes.add_table(11, 7, Inches(0.8), Inches(1.8), Inches(11.733), Inches(5.1))
t10 = table_shape10.table

t10.columns[0].width = Inches(0.8)
t10.columns[1].width = Inches(1.0)
t10.columns[2].width = Inches(1.0)
t10.columns[3].width = Inches(1.6)
t10.columns[4].width = Inches(4.3)
t10.columns[5].width = Inches(1.5)
t10.columns[6].width = Inches(1.533)

headers10 = ["순위", "점수", "자치구", "사건번호", "소재지 및 단지명", "감정평가액", "최저매각가격"]
for col_idx, h in enumerate(headers10):
    cell = t10.cell(0, col_idx)
    cell.text = h
    cell.fill.solid()
    cell.fill.fore_color.rgb = COLOR_NAVY
    p = cell.text_frame.paragraphs[0]
    p.font.size = Pt(10)
    p.font.bold = True
    p.font.color.rgb = COLOR_WHITE
    p.alignment = PP_ALIGN.CENTER

top10_data = [
    ("1위", "100점", "유성구", "2025타경503878", "송강동 청솔아파트 102동 7층", "1억 5,900만", "3,817만 (-76%)"),
    ("2위", "99점", "서구", "2025타경502021", "월평동 에코리움 5층", "1억 8,000만", "3,025만 (-83%)"),
    ("3위", "96점", "유성구", "2025타경503678", "봉명동 사이언스타운 6층 (오피스텔)", "2억", "3,361만 (-83%)"),
    ("4위", "96점", "유성구", "2024타경109647", "봉명동 유성자이 1층 133호 (상가)", "2억 7,600만", "4,638만 (-83%)"),
    ("5위", "96점", "유성구", "2024타경8362", "덕명동 베네디또빌 1층 102호 (다세대)", "4억 1,300만", "9,916만 (-76%)"),
    ("6위", "96점", "유성구", "2024타경8362", "덕명동 베네디또빌 1층 101호 (다세대)", "4억 3,400만", "1억 420만 (-76%)"),
    ("7위", "96점", "유성구", "2025타경3894", "봉명동 다세대 5층 511호", "9,400만", "3,224만 (-66%)"),
    ("8위", "96점", "대덕구", "2025타경503447", "대화동 세원풍남 5동 5층 (아파트)", "1억 6,600만", "5,693만 (-66%)"),
    ("9위", "96점", "유성구", "2026타경45", "용산동 세안프라자 3층 (근린상가)", "1억 7,500만", "6,002만 (-66%)"),
    ("10위", "96점", "유성구", "2025타경3906", "문지동 효성해링턴 배후 상가 3층", "2억 3,600만", "8,094만 (-66%)")
]

for row_idx, r in enumerate(top10_data, start=1):
    for col_idx, val in enumerate(r):
        cell = t10.cell(row_idx, col_idx)
        cell.text = val
        cell.fill.solid()
        cell.fill.fore_color.rgb = COLOR_CARD_BG if row_idx % 2 == 1 else RGBColor(241, 245, 249)
        p = cell.text_frame.paragraphs[0]
        p.font.size = Pt(9.5)
        p.font.color.rgb = COLOR_TEXT_MAIN
        if col_idx in [0, 1, 2, 3]:
            p.alignment = PP_ALIGN.CENTER
            if col_idx == 0 and row_idx <= 3:
                p.font.bold = True
        elif col_idx in [5, 6]:
            p.alignment = PP_ALIGN.RIGHT
            if col_idx == 6:
                p.font.bold = True
                p.font.color.rgb = COLOR_RED

# ==============================================================================
# SLIDE 11: 프로젝트 산출물 및 향후 활용 가이드
# ==============================================================================
slide11 = prs.slides.add_slide(blank_layout)
set_slide_background(slide11, COLOR_GRAY_BG)
add_header(slide11, "프로젝트 산출물 및 실무 활용 가이드", "Project Outputs")

outputs = [
    ("📊 데이터 분석 엑셀 파일 (2종)", "• 대전광역시_경매물건_TOP50_추천.xlsx\n   - 대전 522건 전수 수집 및 100점 만점 TOP 50 정렬\n• 대한민국법원_전국경매물건_정리_및_권역별추천.xlsx\n   - 전국 17개 시·도 1,287건 수집 및 권역별 추천 리스트"),
    ("📄 문서 및 보고서 산출물", "• 부동산_경매_종합_분석_보고서.html\n   - 차트·배너 이미지 4종 탑재, 반응형 웹 스타일\n• 부동산_경매_종합_분석_보고서.md\n   - GitHub 표준 마크다운 보고서\n• 부동산_경매_종합_분석_보고서.pptx\n   - 16:9 고화질 프리젠테이션 (본 파일)"),
    ("⚙️ 자동화 데이터 파이프라인 스크립트", "• process_daejeon_top50.py (대전 전수 수집 및 가치평가)\n• collect_and_export_auction.py (전국 17개 시도 크롤러)\n• generate_charts.py (고해상도 시각화 차트 생성기)"),
    ("💡 실전 입찰 최종 체크리스트", "• 매각기일 1주일 전 매각물건명세서 변경 여부 재확인\n• 해당 아파트/상가 관리사무소 방문해 공용 관리비 확인\n• 네이버 부동산 및 현장 공인중개사를 통한 급매 호가 크로스체크")
]

for idx, (title, desc) in enumerate(outputs):
    col = idx % 2
    row = idx // 2
    x = Inches(0.8 + col * 5.95)
    y = Inches(1.8 + row * 2.5)
    
    card = slide11.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(5.75), Inches(2.25))
    card.fill.solid()
    card.fill.fore_color.rgb = COLOR_CARD_BG
    card.line.color.rgb = COLOR_LINE
    
    tb = slide11.shapes.add_textbox(x + Inches(0.2), y + Inches(0.2), Inches(5.35), Inches(1.9))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = COLOR_BLUE
    
    p = tf.add_paragraph()
    p.text = desc
    p.font.size = Pt(10.5)
    p.font.color.rgb = COLOR_TEXT_MAIN
    p.space_before = Pt(8)

# 저장 실행
output_path1 = r"C:\Users\wood\Desktop\aiwork\auction\부동산_경매_종합_분석_보고서.pptx"
output_path2 = r"C:\Users\wood\Desktop\auction\부동산_경매_종합_분석_보고서.pptx"

prs.save(output_path1)
print(f"PPTX saved successfully: {output_path1}")

try:
    prs.save(output_path2)
    print(f"PPTX backup saved: {output_path2}")
except Exception as e:
    print(f"Backup save skipped: {e}")

